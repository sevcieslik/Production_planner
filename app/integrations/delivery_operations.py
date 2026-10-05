from __future__ import annotations

import os
import re
from datetime import datetime
from io import BytesIO
from typing import Any
from zoneinfo import ZoneInfo

import gspread
import openpyxl

from app.integrations.google_sheets import GoogleSheetsConfigurationError, google_client

DELIVERY_OPERATIONS_ENV = "DELIVERY_OPERATIONS_SPREADSHEET_ID"
RATES_TRACKER_ENV = "RATES_TRACKER_SPREADSHEET_ID"

FLOW_SHEET_NAME = "Flow"
RATES_FRONT_SHEET_NAME = "Front page"


def _configured_id(env_name: str, explicit: str | None = None) -> str:
    value = (explicit or os.getenv(env_name) or "").strip()
    if not value:
        raise GoogleSheetsConfigurationError(f"{env_name} is not configured.")
    return value


def _text(value: Any) -> str:
    return str(value or "").strip()


def _number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = _text(value).replace(",", "")
    if not text or text.upper() in {"N/A", "NA", "-", "NONE"}:
        return None
    match = re.match(r"^\s*([-+]?\d+(?:\.\d+)?)", text)
    if not match:
        return None
    try:
        return float(match.group(1))
    except ValueError:
        return None


def _find_row(values: list[list[Any]], label: str, start: int = 0) -> int | None:
    target = label.strip().lower()
    for index in range(start, len(values)):
        row = values[index]
        first = _text(row[0]) if row else ""
        if first.lower() == target:
            return index
    return None


def _next_nonempty(values: list[list[Any]], start: int) -> int | None:
    for index in range(start, len(values)):
        if any(_text(cell) for cell in values[index]):
            return index
    return None


def _table_from(values: list[list[Any]], header_index: int) -> list[dict[str, Any]]:
    headers = [_text(value) for value in values[header_index]]
    width = len(headers)
    records: list[dict[str, Any]] = []
    for row in values[header_index + 1 :]:
        cells = list(row[:width]) + [""] * max(0, width - len(row))
        if not any(_text(cell) for cell in cells):
            break
        record = {
            header: cells[index]
            for index, header in enumerate(headers)
            if header
        }
        records.append(record)
    return records


def _normalise_flow_records(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    output: list[dict[str, Any]] = []
    for raw in records:
        row: dict[str, Any] = {}
        for key, value in raw.items():
            if key == "Project":
                row[key] = _text(value)
                continue
            numeric = _number(value)
            row[key] = numeric if numeric is not None else _text(value)
        output.append(row)
    return output


def parse_flow_values(values: list[list[Any]]) -> dict[str, Any]:
    """Parse the three visible Flow blocks without relying on fixed row numbers."""
    snapshot = ""
    for row in values[:10]:
        first = _text(row[0]) if row else ""
        if first.lower().startswith("latest source snapshot:"):
            snapshot = first.split(":", 1)[1].strip()
            break

    current_marker = _find_row(values, "Current Production State")
    movement_marker = _find_row(values, "Movement Since Tuesday 15:00")
    backlog_marker = _find_row(values, "Production Backlog")

    current_rows: list[dict[str, Any]] = []
    movement_rows: list[dict[str, Any]] = []
    backlog_rows: list[dict[str, Any]] = []
    movement_period = ""

    if current_marker is not None:
        header = _next_nonempty(values, current_marker + 1)
        if header is not None:
            current_rows = _normalise_flow_records(_table_from(values, header))

    if movement_marker is not None:
        period_row = _next_nonempty(values, movement_marker + 1)
        if period_row is not None:
            movement_period = _text(values[period_row][0]) if values[period_row] else ""
            header = _next_nonempty(values, period_row + 1)
            if header is not None:
                movement_rows = _normalise_flow_records(_table_from(values, header))

    if backlog_marker is not None:
        header = _next_nonempty(values, backlog_marker + 1)
        if header is not None:
            raw_backlog = _table_from(values, header)
            for raw in raw_backlog:
                row: dict[str, Any] = {}
                for key, value in raw.items():
                    if key in {"As Of", "Team", "Status"}:
                        row[key] = _text(value)
                    else:
                        numeric = _number(value)
                        row[key] = numeric if numeric is not None else _text(value)
                backlog_rows.append(row)

    return {
        "snapshot": snapshot,
        "current_state": current_rows,
        "movement_period": movement_period,
        "movement": movement_rows,
        "backlog": backlog_rows,
    }


def read_flow_dashboard(
    *,
    spreadsheet_id: str | None = None,
    client: gspread.Client | None = None,
) -> dict[str, Any]:
    spreadsheet_id = _configured_id(DELIVERY_OPERATIONS_ENV, spreadsheet_id)
    client = client or google_client()
    worksheet = client.open_by_key(spreadsheet_id).worksheet(FLOW_SHEET_NAME)
    return parse_flow_values(worksheet.get_all_values())


def parse_rates_front_page(values: list[list[Any]]) -> list[dict[str, Any]]:
    if not values:
        return []
    headers = [_text(value) for value in values[0]]
    output: list[dict[str, Any]] = []
    for raw in values[1:]:
        cells = list(raw) + [""] * max(0, len(headers) - len(raw))
        project = _text(cells[0]) if cells else ""
        if not project or not project.upper().startswith(("NM", "NA")):
            continue
        record: dict[str, Any] = {}
        for index, header in enumerate(headers):
            if not header:
                continue
            value = cells[index] if index < len(cells) else ""
            if header in {"RS Hours", "GIS Hours", "PLS Hours", "Total"}:
                record[header] = _number(value)
            elif header == "With Richard?":
                record[header] = _text(value).upper() == "TRUE"
            else:
                record[header] = _text(value)
        output.append(record)
    return output


def read_rates_tracker_front_page(
    *,
    spreadsheet_id: str | None = None,
    client: gspread.Client | None = None,
) -> list[dict[str, Any]]:
    spreadsheet_id = _configured_id(RATES_TRACKER_ENV, spreadsheet_id)
    client = client or google_client()
    worksheet = client.open_by_key(spreadsheet_id).worksheet(RATES_FRONT_SHEET_NAME)
    return parse_rates_front_page(worksheet.get_all_values())


def parse_rates_project_values(values: list[list[Any]]) -> dict[str, Any]:
    project_name = ""
    project_code = ""
    for row in values:
        for index, cell in enumerate(row):
            label = _text(cell)
            if label == "Project Name" and index + 1 < len(row):
                project_name = _text(row[index + 1])
            elif label == "Project Code" and index + 1 < len(row):
                project_code = _text(row[index + 1])

    item_header = None
    for index, row in enumerate(values):
        if row and _text(row[0]).lower() == "item":
            item_header = index
            break

    assumptions: list[dict[str, Any]] = []
    if item_header is not None:
        headers = [_text(value) for value in values[item_header][:5]]
        for raw in values[item_header + 1 :]:
            cells = list(raw[:5]) + [""] * max(0, 5 - len(raw))
            item = _text(cells[0])
            if not item:
                break
            assumptions.append(
                {
                    headers[0] or "Item": item,
                    headers[1] or "Rate km/day": _number(cells[1]),
                    headers[2] or "Rate Per km": _text(cells[2]),
                    headers[3] or "Team": _text(cells[3]),
                    headers[4] or "Hrs": _number(cells[4]),
                }
            )

    return {
        "project_name": project_name,
        "project_code": project_code,
        "assumptions": assumptions,
    }


def read_rates_tracker_project(
    project_code: str,
    *,
    spreadsheet_id: str | None = None,
    client: gspread.Client | None = None,
) -> dict[str, Any]:
    spreadsheet_id = _configured_id(RATES_TRACKER_ENV, spreadsheet_id)
    client = client or google_client()
    spreadsheet = client.open_by_key(spreadsheet_id)
    code = _text(project_code).upper()
    worksheet = next(
        (
            ws for ws in spreadsheet.worksheets()
            if ws.title.strip().upper().startswith(code)
        ),
        None,
    )
    if worksheet is None:
        return {
            "project_name": "",
            "project_code": code,
            "assumptions": [],
            "sheet_title": None,
        }
    parsed = parse_rates_project_values(worksheet.get_all_values())
    parsed["sheet_title"] = worksheet.title
    return parsed


LEGACY_HISTORY_COLUMNS = {
    "TPC Outstanding km": "tpc_outstanding_km",
    "TPC WIP km": "tpc_wip_km",
    "GIS Ready km": "gis_ready_km",
    "GIS Production WIP km": "gis_production_wip_km",
    "GIS Subcon WIP km": "gis_subcon_wip_km",
    "GIS Ready For QC km": "gis_ready_for_qc_km",
    "GIS QC WIP km": "gis_qc_wip_km",
    "PLS Ready km": "pls_ready_km",
    "PLS WIP km": "pls_wip_km",
}

LEGACY_BACKLOG_COLUMNS = {
    "Queue km": "queue_km",
    "Historical Capacity km/week": "historical_capacity_km_week",
    "Backlog Weeks": "backlog_weeks",
    "Provisional Backlog Weeks": "provisional_backlog_weeks",
    "Weeks Used": "weeks_used",
    "Lookback Weeks": "lookback_weeks",
    "Target Weeks": "target_weeks",
    "Status": "status",
}

_LONDON = ZoneInfo("Europe/London")


def _aware_datetime(value: Any) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value if value.tzinfo is not None else value.replace(tzinfo=_LONDON)
    text = _text(value)
    for fmt in ("%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).replace(tzinfo=_LONDON)
        except ValueError:
            pass
    return None


def _worksheet_values(workbook, title: str) -> list[list[Any]]:
    if title not in workbook.sheetnames:
        return []
    worksheet = workbook[title]
    values: list[list[Any]] = []
    for row in worksheet.iter_rows(values_only=True):
        values.append(list(row))
    return values


def _records_from_sheet(values: list[list[Any]]) -> list[dict[str, Any]]:
    if not values:
        return []
    headers = [_text(value) for value in values[0]]
    output: list[dict[str, Any]] = []
    for raw in values[1:]:
        if not any(value not in (None, "") for value in raw):
            continue
        row = {}
        for index, header in enumerate(headers):
            if not header:
                continue
            row[header] = raw[index] if index < len(raw) else None
        output.append(row)
    return output


def parse_delivery_operations_workbook(content: bytes) -> dict[str, Any]:
    """Parse a legacy Delivery Operations workbook for one-time DB migration.

    This is deliberately a file parser, not a live Google Sheets dependency.
    """
    workbook = openpyxl.load_workbook(BytesIO(content), data_only=True, read_only=False)

    flow_values = _worksheet_values(workbook, "Flow")
    dashboard = parse_flow_values(flow_values) if flow_values else None
    dashboard_snapshot_at = None
    if dashboard and dashboard.get("snapshot"):
        dashboard_snapshot_at = _aware_datetime(dashboard["snapshot"])

    history: list[dict[str, Any]] = []
    for raw in _records_from_sheet(_worksheet_values(workbook, "_tracker_history")):
        snapshot_at = _aware_datetime(raw.get("Snapshot Date"))
        project_code = _text(raw.get("Project Code"))
        project_name = _text(raw.get("Project Name"))
        if snapshot_at is None or not project_code or not project_name:
            continue
        metrics = {}
        for legacy_header, key in LEGACY_HISTORY_COLUMNS.items():
            value = _number(raw.get(legacy_header))
            metrics[key] = 0.0 if value is None else value
        history.append(
            {
                "snapshot_at": snapshot_at,
                "project_code": project_code,
                "project_name": project_name,
                "metrics": metrics,
            }
        )

    movements: list[dict[str, Any]] = []
    movement_rows = _records_from_sheet(_worksheet_values(workbook, "_tracker_movements"))
    fixed = {"from_snapshot", "to_snapshot", "project_code", "project_name", "portfolio_status"}
    for raw in movement_rows:
        from_snapshot = _aware_datetime(raw.get("from_snapshot"))
        to_snapshot = _aware_datetime(raw.get("to_snapshot"))
        project_code = _text(raw.get("project_code"))
        project_name = _text(raw.get("project_name"))
        if from_snapshot is None or to_snapshot is None or not project_code or not project_name:
            continue
        metrics = {}
        for key, value in raw.items():
            if key in fixed or not key:
                continue
            number = _number(value)
            metrics[key] = 0.0 if number is None else number
        movements.append(
            {
                "from_snapshot": from_snapshot,
                "to_snapshot": to_snapshot,
                "project_code": project_code,
                "project_name": project_name,
                "portfolio_status": _text(raw.get("portfolio_status")) or None,
                "metrics": metrics,
            }
        )

    backlog: list[dict[str, Any]] = []
    for raw in _records_from_sheet(_worksheet_values(workbook, "_flow_kpis")):
        as_of = _aware_datetime(raw.get("As Of"))
        team = _text(raw.get("Team"))
        if as_of is None or not team:
            continue
        metrics: dict[str, Any] = {}
        for legacy_header, key in LEGACY_BACKLOG_COLUMNS.items():
            value = raw.get(legacy_header)
            if key == "status":
                metrics[key] = _text(value)
            else:
                metrics[key] = _number(value)
        backlog.append({"as_of": as_of, "team": team, "metrics": metrics})

    return {
        "dashboard": dashboard,
        "dashboard_snapshot_at": dashboard_snapshot_at,
        "history": history,
        "movements": movements,
        "backlog": backlog,
    }
